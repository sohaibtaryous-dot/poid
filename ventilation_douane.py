# -*- coding: utf-8 -*-
"""Ventilation douane — Alstom Railways Maroc (Fès) : un classeur Excel par expédition / fournisseur, deux feuilles :
  « Ventilation »  : une ligne par ligne de facture (jamais de fusion de références), formules vivantes
                     (poids net g × qté, kg, HT, frais ventilés au prorata, total HT + frais), NGP du référentiel, couleurs par facture ;
  « Contrôle BL »  : une ligne par facture (SUMIF sur Ventilation), écarts net / brut / HT, rapprochement avec l'avis d'arrivée
                     (manifeste : valeur et poids), notes et anomalies (« Ntebh »).
Plus un récapitulatif consolidé (Détail / Total par Fournisseur / Total Global).
Règles : poids manquants estimés par prorata (marqués en jaune foncé + commentaire), frais administratifs ventilés au prorata du HT
(dernière ligne = reste), NGP absent du référentiel = « A VERIFIER », NGP ≠ code HS de la facture = anomalie signalée."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import import_checks as ic
import ventilation_export as vx
import ventilation_import as vi

NAVY, ORANGE, EST = "1F3864", "C65911", "7F6000"
PALETTE = ["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "EDEDED", "E4DFEC", "DAEEF3", "F8CBAD"]
COLS = ["N°", "Expédition", "N° BL", "N° Facture", "Poste", "Référence", "Désignation", "HS Code", "COO", "Unité", "Qté",
        "Poids net unitaire (g)", "Poids net total ligne (g)", "Poids net total ligne (kg)", "Poids brut ligne (kg)",
        "Prix unitaire ({cur})", "Total HT ligne ({cur})", "Frais administratifs ventilés ({cur})", "Total HT + frais ({cur})", "NGP"]
WIDTHS = [5, 16, 14, 18, 8, 22, 36, 14, 7, 7, 10, 14, 14, 14, 14, 12, 14, 14, 14, 14]
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _is_fee(v) -> bool:
    if isinstance(v, str):
        return v.strip().lower() in ("true", "1", "oui", "yes", "vrai")
    return bool(v) if v is not None and v == v else False


def _num(v, default=0.0) -> float:
    x = pd.to_numeric(pd.Series([v]), errors="coerce").iloc[0]
    return float(x) if x == x else default


def _colis_txt(v) -> str:
    x = _num(v, None)
    if x is None:
        t = str(v if v is not None else "").strip()
        return "" if t.lower() in ("nan", "none") else t
    return f"{x:g}"


def safe_name(s: str) -> str:
    return re.sub(r"[^\w\-]+", "_", str(s or "SANS_NOM")).strip("_")[:40] or "SANS_NOM"


# ----------------------------------------------------------------------------- préparation des lignes d'un fournisseur
def prepare(final: pd.DataFrame, suppliers: list[str], ratio: float | None = None):
    """Lignes de marchandises (dicts) + frais par facture, pour un ou plusieurs fournisseurs mis dans le même tableau."""
    d = final.copy().reset_index(drop=True)
    d["Client"] = d["Client"].map(lambda v: str(v or "").strip() or "(sans nom)")
    sub = d[d["Client"].isin(suppliers)].reset_index(drop=True)
    notes: list[str] = []
    dummy = vx.Reference({}, {}, {}, {})
    norm, _tot, n_notes = vi.normalize(sub, dummy, ratio)               # poids net (estimé si absent), lignes de frais exclues
    notes.extend(n for n in n_notes if "estimé" in n or "sans poids" in n)
    fee_mask = sub["_fee"].map(_is_fee) if "_fee" in sub else pd.Series([False] * len(sub))
    if fee_mask.all():
        fee_mask = pd.Series([False] * len(sub))
    goods = sub[~fee_mask.values].reset_index(drop=True)
    fees = sub[fee_mask.values]
    fee_by_inv: dict[str, float] = {}
    for _, r in fees.iterrows():
        k = str(r.get("N° facture") or "").strip()
        fee_by_inv[k] = fee_by_inv.get(k, 0.0) + _num(r.get("Valeur HT"))
    total_by_inv = sub.assign(_k=sub["N° facture"].map(lambda v: str(v or "").strip())).groupby("_k")["Valeur HT"].apply(
        lambda s: float(pd.to_numeric(s, errors="coerce").fillna(0).sum())).to_dict()
    lines = []
    for i, (_, r) in enumerate(goods.iterrows()):
        n = norm.iloc[i] if i < len(norm) else None
        qty = _num(r.get("Qté"))
        net_kg = float(n["_net"]) if n is not None else _num(r.get(vi.C_NET))
        lines.append({
            "supplier": r["Client"], "bl": str(r.get("_bl") or "").strip() or "—", "exp": str(r.get("Job Ref manifest") or "").strip(), "inv": str(r.get("N° facture") or "").strip() or "(sans facture)",
            "pos": (lambda v: v.lstrip("0") or "0" if v.isdigit() else v)(str(r.get("Item") or "").strip()), "ref": str(r.get("Réf article") or "").strip(), "desc": str(r.get("Désignation") or "").strip(),
            "hs": str(r.get("NGP/code douanier (facture)") or "").strip(), "coo": str(r.get("Origine") or "").strip(), "qty": qty,
            "net_kg": net_kg, "net_est": bool(n["_net_est"]) if n is not None else False,
            "price": _num(r.get("Prix unitaire"), None), "ht": _num(r.get("Valeur HT"), None),
            "gross_inv": _num(r.get(vi.C_GROSSF)), "netf_inv": _num(r.get(vi.C_NETF)), "colis": str(r.get(vi.C_COLIS) or "").strip(),
            "ngp_ref": str(r.get("NGP (REF_PAR_FRS)") or "").strip(), "ctrl": str(r.get("Contrôle NGP") or "").strip(),
            "cur": str(r.get("_devise") or "").strip(), "src": str(r.get("_source") or "").strip(), "remark": str(r.get("Remarque") or "").strip()})
    return lines, fee_by_inv, total_by_inv, notes


def _avis(man: pd.DataFrame | None, suppliers: list[str]) -> dict:
    """Avis d'arrivée / manifeste du ou des fournisseurs : expéditions, valeur, poids, colis."""
    if man is None or man.empty or "Expéditeur" not in man:
        return {}
    names = [str(x) for x in man["Expéditeur"]]
    rows = pd.DataFrame()
    for s in suppliers:
        m = ic.match_name(s, names)
        if m and not m.upper().startswith("TOTAL"):
            rows = pd.concat([rows, man[man["Expéditeur"].astype(str) == m]])
    if rows.empty:
        return {}
    g = lambda c: float(pd.to_numeric(rows[c], errors="coerce").sum()) if c in rows else None  # noqa: E731
    return {"value": g("Goods Value"), "brut": g("Poids (kg)"), "colis": g("Colis"),
            "ship": ", ".join(dict.fromkeys(str(x) for x in rows.get("Shipment", []) if str(x) not in ("", "nan", "None")))}


# ----------------------------------------------------------------------------- classeur d'un fournisseur
def build_supplier_workbook(final: pd.DataFrame, suppliers: list[str], man: pd.DataFrame | None = None, ratio: float | None = None,
                            title: str | None = None) -> tuple[bytes, dict]:
    lines, fee_by_inv, total_by_inv, notes = prepare(final, suppliers, ratio)
    cur = next((l["cur"] for l in lines if l["cur"]), "EUR")
    avis = _avis(man, suppliers)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventilation"
    hdr = [c.format(cur=cur) for c in COLS]
    ws.append(hdr)
    for j, c in enumerate(ws[1], start=1):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=ORANGE if hdr[j - 1].startswith(("Frais", "Total HT + frais")) else NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 42
    keys = list(dict.fromkeys((l["bl"], l["inv"]) for l in lines))            # un bloc par BL / facture
    by_key = {k: [l for l in lines if (l["bl"], l["inv"]) == k] for k in keys}
    invoices = list(dict.fromkeys(l["inv"] for l in lines))
    by_inv = {inv: [l for l in lines if l["inv"] == inv] for inv in invoices}
    fee_used: set = set()
    first_row = {}
    r = 2
    est_net = any(l["net_est"] for l in lines)
    est_brut = False
    dup_refs = sorted({l["ref"] for l in lines if l["ref"] and sum(1 for x in lines if x["ref"] == l["ref"]) > 1})
    mismatch, missing_ngp = [], 0
    n, first_est = 0, False
    for b, (bl_k, inv) in enumerate(keys):
        fill = PatternFill("solid", start_color=PALETTE[b % len(PALETTE)])
        grp = by_key[(bl_k, inv)]
        r0, r1 = r, r + len(grp) - 1
        fee = fee_by_inv.get(inv, 0.0) if inv not in fee_used else 0.0
        fee_used.add(inv)
        ht_vals = [(l["ht"] if l["ht"] is not None else (round((l["qty"] or 0) * (l["price"] or 0), 2))) for l in grp]
        sum_ht = sum(ht_vals) or 1.0
        gross_inv = grp[0]["gross_inv"]
        first_row[inv] = r0
        for k, l in enumerate(grp):
            n += 1
            ws.cell(r, 1, n)
            ws.cell(r, 2, l["exp"] or (avis.get("ship") or ""))
            ws.cell(r, 3, bl_k)
            ws.cell(r, 4, inv)
            ws.cell(r, 5, l["pos"])
            ws.cell(r, 6, l["ref"])
            ws.cell(r, 7, l["desc"])
            ws.cell(r, 8, l["hs"])
            ws.cell(r, 9, l["coo"])
            ws.cell(r, 10, "NB")
            ws.cell(r, 11, l["qty"])
            ws.cell(r, 12, (l["net_kg"] * 1000.0 / l["qty"]) if l["qty"] else 0.0)           # g / pièce
            ws.cell(r, 13, f"=K{r}*L{r}")
            ws.cell(r, 14, f"=M{r}/1000")
            brut = (gross_inv * ht_vals[k] / sum_ht) if gross_inv else 0.0                    # brut de la facture réparti au prorata du HT
            est_brut = est_brut or bool(gross_inv)
            ws.cell(r, 15, round(brut, 3))
            ws.cell(r, 16, l["price"] if l["price"] is not None else None)
            if l["ht"] is not None:
                ws.cell(r, 17, round(l["ht"], 2))
            else:
                ws.cell(r, 17, f"=ROUND(K{r}*P{r},2)")
            if fee and len(grp) > 1 and k == len(grp) - 1:
                ws.cell(r, 18, f"={round(fee, 2)}-SUM(R{r0}:R{r - 1})")                      # la dernière ligne absorbe l'arrondi
            elif fee:
                ws.cell(r, 18, f"=ROUND(Q{r}/SUM(Q${r0}:Q${r1})*{round(fee, 2)},2)")
            else:
                ws.cell(r, 18, 0)
            ws.cell(r, 19, f"=Q{r}+R{r}")
            ngp = vi.ngp_digits(l["ngp_ref"])
            ws.cell(r, 20, ngp or "A VERIFIER")
            if not ngp:
                missing_ngp += 1
                ws.cell(r, 20).font = Font(bold=True, color="C00000")
            elif l["ctrl"] == "ÉCART":
                mismatch.append((l["ref"], l["hs"], ngp))
                ws.cell(r, 20).comment = Comment(f"Ntebh : NGP du référentiel {ngp} ≠ code HS de la facture {l['hs']}", "TIMAR")
                ws.cell(r, 20).font = Font(bold=True, color="C00000")
            if l["net_est"] and not first_est:
                ws.cell(r, 12).comment = Comment("Poids net estimé (prorata) : voir le commentaire de l'en-tête.", "TIMAR")
                first_est = True
            for c in range(1, 21):
                cell = ws.cell(r, c)
                cell.fill = fill
                cell.border = BORDER
                if c in (12, 13, 14, 15, 16, 17, 18, 19):
                    cell.number_format = "#,##0.000" if c in (12, 14, 15) else "#,##0.00"
                if c == 11:
                    cell.number_format = "#,##0.###"
            r += 1
    last = max(r - 1, 2)
    tot = r
    ws.cell(tot, 7, "TOTAL").font = Font(bold=True)
    for c in (11, 13, 14, 15, 17, 18, 19):
        col = get_column_letter(c)
        cell = ws.cell(tot, c, f"=SUBTOTAL(9,{col}2:{col}{last})")
        cell.font = Font(bold=True)
        cell.number_format = "#,##0.000" if c in (14, 15) else "#,##0.00"
    for c in range(1, 21):
        ws.cell(tot, c).border = Border(top=Side(style="medium"), bottom=Side(style="medium"))
    # poids estimés : en-têtes jaune foncé + commentaire
    def mark(col_idx: int, text: str):
        cell = ws.cell(1, col_idx)
        cell.fill = PatternFill("solid", start_color=EST)
        cell.comment = Comment(text, "TIMAR")
    if est_net:
        ratio_txt = f"{(ratio if ratio is not None else vi.default_ratio(final)) * 100:.0f} %"
        for ci in (12, 13, 14):
            mark(ci, f"Poids net estimé pour une partie des lignes : poids net de la facture, sinon poids brut × {ratio_txt}, réparti au prorata de la quantité.")
    if est_brut:
        mark(15, "Poids brut de la facture réparti sur ses lignes au prorata de la valeur HT (pas de poids brut par ligne sur les documents).")
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:T{last}"
    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ------------------------------------------------------------------ Contrôle BL
    cs = wb.create_sheet("Contrôle BL")
    ch = ["Expédition", "N° BL", "N° Facture", "Nb lignes", "Qté totale", "Poids net calculé (kg)", "Poids net déclaré (kg)", "Écart net (kg)",
          "Poids brut lignes (kg)", "Poids brut déclaré (kg)", "Écart brut (kg)", "Colis", f"Total HT marchandises ({cur})",
          f"Frais administratifs ({cur})", f"Total HT + frais calculé ({cur})", f"Total HT facture ({cur})", f"Écart HT ({cur})"]
    cs.append(ch)
    for c in cs[1]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color=NAVY), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cs.row_dimensions[1].height = 42
    rng = lambda col: f"Ventilation!${col}$2:${col}${last}"  # noqa: E731
    fee_seen: set = set()
    for b, (bl_k, inv) in enumerate(keys):
        rr = b + 2
        grp = by_key[(bl_k, inv)]
        netf, gross = grp[0]["netf_inv"], grp[0]["gross_inv"]
        cond = f"{rng('C')},B{rr},{rng('D')},C{rr}"
        cs.cell(rr, 1, grp[0]["exp"] or (avis.get("ship") or ""))
        cs.cell(rr, 2, bl_k)
        cs.cell(rr, 3, inv)
        cs.cell(rr, 4, f"=COUNTIFS({cond})")
        cs.cell(rr, 5, f"=SUMIFS({rng('K')},{cond})")
        cs.cell(rr, 6, f"=SUMIFS({rng('N')},{cond})")
        cs.cell(rr, 7, netf if netf else None)
        cs.cell(rr, 8, f'=IF(G{rr}="","",ROUND(F{rr}-G{rr},3))')
        cs.cell(rr, 9, f"=SUMIFS({rng('O')},{cond})")
        cs.cell(rr, 10, gross if gross else None)
        cs.cell(rr, 11, f'=IF(J{rr}="","",ROUND(I{rr}-J{rr},3))')
        cs.cell(rr, 12, _colis_txt(grp[0]["colis"]))
        cs.cell(rr, 13, f"=SUMIFS({rng('Q')},{cond})")
        cs.cell(rr, 14, f"=SUMIFS({rng('R')},{cond})")
        cs.cell(rr, 15, f"=SUMIFS({rng('S')},{cond})")
        cs.cell(rr, 16, round(total_by_inv.get(inv, 0.0), 2) if inv not in fee_seen else 0.0)
        fee_seen.add(inv)
        cs.cell(rr, 17, f"=ROUND(O{rr}-P{rr},2)")
        for c in range(1, 18):
            cs.cell(rr, c).fill = PatternFill("solid", start_color=PALETTE[b % len(PALETTE)])
            cs.cell(rr, c).border = BORDER
            if c >= 5 and c != 12:
                cs.cell(rr, c).number_format = "#,##0.000" if c in (6, 7, 8, 9, 10, 11) else "#,##0.00"
    t = len(keys) + 2
    cs.cell(t, 1, "TOTAL").font = Font(bold=True)
    for c in (4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17):
        col = get_column_letter(c)
        cell = cs.cell(t, c, f"=SUM({col}2:{col}{t - 1})")
        cell.font = Font(bold=True)
        cell.number_format = "#,##0.000" if c in (6, 7, 8, 9, 10, 11) else "#,##0.00"
    colis_n = sum(_num(g[0]["colis"]) for g in by_key.values())
    cs.cell(t, 12, f"{colis_n:g} colis" if colis_n else "")
    # rapprochement avec l'avis d'arrivée (libellés en A, valeurs en B — comme dans le tableau de référence)
    rr = t + 2
    cs.cell(rr, 1, "Poids brut déclaré avis d'arrivée (kg)")
    cs.cell(rr, 2, avis.get("brut"))
    cs.cell(rr + 1, 1, "Écart poids brut (lignes − avis)")
    cs.cell(rr + 1, 2, f'=IF(B{rr}="","",ROUND(I{t}-B{rr},3))')
    cs.cell(rr + 2, 1, f"Valeur avis d'arrivée (GOODS VALUE, {cur})")
    cs.cell(rr + 2, 2, avis.get("value"))
    cs.cell(rr + 3, 1, "Écart valeur (Total HT + frais − avis)")
    cs.cell(rr + 3, 2, f'=IF(B{rr + 2}="","",ROUND(O{t}-B{rr + 2},2))')
    cs.cell(rr + 4, 1, "Colis avis d'arrivée")
    cs.cell(rr + 4, 2, avis.get("colis"))
    for k in range(5):
        cs.cell(rr + k, 1).font = Font(bold=True)
        cs.cell(rr + k, 2).number_format = "#,##0.00"
    for k in (1, 3):
        cs.cell(rr + k, 2).font = Font(bold=True)
    # notes
    srcs = sorted({l["src"] for l in lines if l["src"]})
    cur_set = sorted({l["cur"] for l in lines if l["cur"]})
    note_rows: list[tuple[str, bool]] = [
        (f"Fournisseur(s) : {', '.join(suppliers)} · devise {', '.join(cur_set) or cur} · {len(lines)} ligne(s), {len(invoices)} facture(s)"
         + (f" · expédition(s) {avis['ship']}" if avis.get("ship") else ""), False),
        ("Documents sources : " + (", ".join(srcs[:6]) + (" …" if len(srcs) > 6 else "") if srcs else "factures lues par le site") + ".", False),
        ("Poids : net lu sur la facture / packing list quand il existe ; sinon estimé (net facture, ou brut × ratio) au prorata de la quantité ; "
         "brut de facture réparti au prorata du HT. Frais administratifs ventilés au prorata du HT (dernière ligne = reste). "
         "Références répétées conservées sur des lignes séparées.", False)]
    flags: list[str] = []
    if est_net:
        flags.append("Poids net estimé sur certaines lignes (en-têtes jaune foncé) : à vérifier avec les documents.")
    if missing_ngp:
        flags.append(f"{missing_ngp} ligne(s) sans NGP dans le référentiel : « A VERIFIER ». Vérifiez que vous avez la dernière version du référentiel (CONSOLIDE_CLIENTS).")
    if mismatch:
        flags.append(f"Ntebh : {len(mismatch)} ligne(s) où le NGP du référentiel diffère du code HS de la facture (le NGP du référentiel est retenu) : "
                     + ", ".join(f"{a} ({b}→{c})" for a, b, c in mismatch[:4]) + (" …" if len(mismatch) > 4 else ""))
    if dup_refs:
        flags.append(f"Références répétées (lignes gardées séparées) : {', '.join(dup_refs[:6])}" + (" …" if len(dup_refs) > 6 else ""))
    if len(cur_set) > 1:
        flags.append(f"Plusieurs devises ({', '.join(cur_set)}) : les totaux additionnent des montants de devises différentes.")
    if not avis:
        flags.append("Aucun avis d'arrivée / manifeste rattaché à ce fournisseur : valeur et poids non rapprochés.")
    note_rows += [(f, True) for f in flags]
    for k, (txt, red) in enumerate(note_rows[:8]):
        c = cs.cell(rr + 6 + k, 1, txt)
        c.font = Font(italic=True, color="C00000" if red else "7F7F7F")
    for j, w in enumerate([40, 14, 14, 9, 12, 13, 13, 10, 13, 13, 10, 24, 15, 13, 15, 15, 11], start=1):
        cs.column_dimensions[get_column_letter(j)].width = w
    cs.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    stats = {"supplier": " + ".join(suppliers), "lignes": len(lines), "factures": len(invoices), "qté": sum(l["qty"] for l in lines),
             "net": sum(l["net_kg"] for l in lines), "brut": sum(g[0]["gross_inv"] for g in by_inv.values()),
             "ht": sum((l["ht"] if l["ht"] is not None else round(l["qty"] * (l["price"] or 0), 2)) for l in lines),
             "frais": sum(fee_by_inv.get(i, 0.0) for i in invoices), "devise": cur, "avis_valeur": avis.get("value"), "avis_brut": avis.get("brut"),
             "A VERIFIER": missing_ngp, "écarts NGP/HS": len(mismatch), "net estimé": sum(1 for l in lines if l["net_est"]), "flags": flags,
             "_lines": lines}
    return out.getvalue(), stats


# ----------------------------------------------------------------------------- récapitulatif consolidé
def build_recap(all_stats: list[dict]) -> bytes:
    wb = Workbook()
    d = wb.active
    d.title = "Détail"
    head = ["Fournisseur"] + [c.format(cur="devise") for c in COLS]
    d.append(head)
    for c in d[1]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color=NAVY), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    d.row_dimensions[1].height = 42
    r, n = 2, 0
    for si, st in enumerate(all_stats):
        fill = PatternFill("solid", start_color=PALETTE[si % len(PALETTE)])
        for l in st["_lines"]:
            n += 1
            ht = l["ht"] if l["ht"] is not None else round(l["qty"] * (l["price"] or 0), 2)
            row = [st["supplier"], n, l["exp"], "", l["inv"], l["pos"], l["ref"], l["desc"], l["hs"], l["coo"], "NB", l["qty"],
                   (l["net_kg"] * 1000 / l["qty"]) if l["qty"] else 0.0, l["net_kg"] * 1000, l["net_kg"], None, l["price"], ht, None, ht,
                   vi.ngp_digits(l["ngp_ref"]) or "A VERIFIER"]
            for j, v in enumerate(row, start=1):
                c = d.cell(r, j, v)
                c.fill, c.border = fill, BORDER
            r += 1
    last = max(r - 1, 2)
    d.cell(r, 8, "TOTAL").font = Font(bold=True)
    for c in (12, 14, 15, 18, 20):
        col = get_column_letter(c)
        d.cell(r, c, f"=SUBTOTAL(9,{col}2:{col}{last})").font = Font(bold=True)
    d.freeze_panes = "I2"
    d.auto_filter.ref = f"A1:U{last}"
    for j, w in enumerate([22] + WIDTHS, start=1):
        d.column_dimensions[get_column_letter(j)].width = w

    t = wb.create_sheet("Total par Fournisseur")
    th = ["Fournisseur", "Lignes", "Qté", "Poids net (kg)", "Poids brut (kg)", "Total HT", "Frais", "Total HT + frais", "Valeur avis d'arrivée", "Écart valeur",
          "A VERIFIER (NGP)", "Écarts NGP/HS"]
    t.append(th)
    for c in t[1]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color=NAVY), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, st in enumerate(all_stats, start=2):
        s = f"'Détail'!$A$2:$A${last}"
        t.cell(i, 1, st["supplier"])
        t.cell(i, 2, f"=COUNTIF({s},A{i})")
        t.cell(i, 3, f"=SUMIF({s},A{i},'Détail'!$L$2:$L${last})")
        t.cell(i, 4, f"=SUMIF({s},A{i},'Détail'!$O$2:$O${last})")
        t.cell(i, 5, st["brut"] or None)
        t.cell(i, 6, f"=SUMIF({s},A{i},'Détail'!$R$2:$R${last})")
        t.cell(i, 7, st["frais"])
        t.cell(i, 8, f"=F{i}+G{i}")
        t.cell(i, 9, st["avis_valeur"])
        t.cell(i, 10, f'=IF(I{i}="","",ROUND(H{i}-I{i},2))')
        t.cell(i, 11, st["A VERIFIER"])
        t.cell(i, 12, st["écarts NGP/HS"])
        for c in range(1, 13):
            t.cell(i, c).fill = PatternFill("solid", start_color=PALETTE[(i - 2) % len(PALETTE)])
            t.cell(i, c).border = BORDER
    tt = len(all_stats) + 2
    t.cell(tt, 1, "TOTAL").font = Font(bold=True)
    for c in range(2, 13):
        col = get_column_letter(c)
        t.cell(tt, c, f"=SUM({col}2:{col}{tt - 1})").font = Font(bold=True)
    for j, w in enumerate([30, 8, 12, 14, 14, 16, 12, 16, 18, 14, 14, 14], start=1):
        t.column_dimensions[get_column_letter(j)].width = w

    g = wb.create_sheet("Total Global")
    g["A1"] = "Total Global du dossier"
    g["A1"].font = Font(bold=True, size=14)
    kp = [("Fournisseurs", f"=COUNTA('Total par Fournisseur'!A2:A{tt - 1})"), ("Lignes", f"='Total par Fournisseur'!B{tt}"),
          ("Quantité totale", f"='Total par Fournisseur'!C{tt}"), ("Poids net (kg)", f"='Total par Fournisseur'!D{tt}"),
          ("Poids brut (kg)", f"='Total par Fournisseur'!E{tt}"), ("Total HT", f"='Total par Fournisseur'!F{tt}"),
          ("Frais administratifs", f"='Total par Fournisseur'!G{tt}"), ("Total HT + frais", f"='Total par Fournisseur'!H{tt}"),
          ("Valeur avis d'arrivée (fournisseurs rattachés)", f"='Total par Fournisseur'!I{tt}"),
          ("Écart valeur", f"='Total par Fournisseur'!J{tt}")]
    for k, (lab, f) in enumerate(kp, start=3):
        g.cell(k, 1, lab).font = Font(bold=True)
        g.cell(k, 2, f).number_format = "#,##0.00"
    g.column_dimensions["A"].width = 46
    g.column_dimensions["B"].width = 20
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_all(final: pd.DataFrame, man: pd.DataFrame | None = None, ratio: float | None = None,
              groups: list[list[str]] | None = None) -> dict:
    """Un classeur par fournisseur (ou par groupe fusionné à la demande) + récapitulatif consolidé + ZIP."""
    names = list(dict.fromkeys(str(v or "").strip() or "(sans nom)" for v in final["Client"]))
    groups = [g for g in (groups or []) if len(g) > 1]
    used = {s for g in groups for s in g}
    units = groups + [[n] for n in names if n not in used]
    files, stats = {}, []
    for u in units:
        data, st = build_supplier_workbook(final, u, man, ratio)
        fname = f"VENTILATION_{safe_name(' + '.join(u))}.xlsx"
        files[fname] = data
        st["file"] = fname
        stats.append(st)
    recap = build_recap(stats) if stats else b""
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items():
            z.writestr(k, v)
        if recap:
            z.writestr("RECAP_CONSOLIDE.xlsx", recap)
    summary = pd.DataFrame([{k: v for k, v in s.items() if k not in ("_lines", "flags")} for s in stats])
    return {"files": files, "recap": recap, "zip": bio.getvalue(), "summary": summary, "stats": stats,
            "stamp": datetime.now().strftime("%Y%m%d_%H%M")}
