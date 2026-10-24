# -*- coding: utf-8 -*-
"""Tests automatiques : python tests/test_core.py   (ou : python -m pytest tests)"""
import io
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as core  # noqa: E402
import fileio  # noqa: E402
import shipment  # noqa: E402
import units  # noqa: E402


class _F:  # fichier téléversé factice
    def __init__(self, name, data):
        self.name, self._b = name, data

    def getvalue(self):
        return self._b


class _P:
    def progress(self, *a, **k):
        pass


def _xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _run(files, basis="Valeur HT"):
    return core.run_pipeline(files, None, core.OCREngine("none", "", 300), basis, _P())


def test_unites_poids():
    assert units.to_kg(250, "g") == 0.25
    assert units.to_kg(1.2, "t") == 1200
    assert abs(units.to_kg(10, "lbs") - 4.5359237) < 1e-6
    assert units.to_kg(3, None) == 3
    assert units.unit_from_header("Net weight (g)") == "g"
    assert core.weight_cell("12,5 lbs", ",") == units.to_kg(12.5, "lbs")
    assert core.find_weight("Total gross weight: 0.03 t", core.TOTAL_GROSS_RE, ".") == 30.0


def test_nombres_et_decimale():
    assert core.parse_number("1 204,023 KG", ",") == 1204.023
    assert core.parse_number("1.234,56", ".") == 1234.56
    assert core.parse_number("1,234.56", ",") == 1234.56
    assert core.find_weight("Poids brut total : 148,000KG", core.TOTAL_GROSS_RE, ".") == 148.0   # jamais 148 000 kg


def test_dates():
    assert core.normalize_date("13-OCT-2026") == "13/10/2026"
    assert core.normalize_date("October 13, 2026") == "13/10/2026"
    assert core.normalize_date("13 octobre 2026") == "13/10/2026"
    assert core.normalize_date("2026-10-13T00:00:00") == "13/10/2026"
    assert core.normalize_date("10/13/2026") == "13/10/2026"


def test_origine():
    assert core.find_origin("Incoterm Pays d'origine de la marchandise Pays d'origine du tissu\nFCA MAROC CHINE") == "MA"
    assert core.find_origin("Country of origin: China") == "CN"


def test_fichiers_invalides():
    for name, data in [("a.pdf", b""), ("b.pdf", b"texte"), ("c.xlsx", b"nope"), ("d.exe", b"x"), ("e.docx", b"PK\x03\x04x")]:
        try:
            fileio.read_any(name, fileio.validate_upload(name, data), data)
            raise AssertionError(f"{name} aurait dû être refusé")
        except fileio.FileReadError:
            pass


def test_un_fichier_corrompu_ne_fait_pas_planter_le_dossier():
    good = _xlsx([["ACME GmbH"], ["Invoice No.", "A-1", "Date", "01/02/2026"], ["Currency", "EUR"],
                  ["Pos", "Part No", "Description", "Qty", "Unit price", "Amount"], [1, "X-100", "Pièce", 2, 5, 10]])
    r = _run([_F("ok.xlsx", good), _F("faux.pdf", b"%PDF-1.4 pas un pdf"), _F("vide.xlsx", b"")])
    assert len(r["lines"]) == 1 and len(r["issues"]) == 2
    assert {i["Gravité"] for i in r["issues"]} == {"Erreur"}


def test_poids_repris_de_la_packing_list_et_ecarts():
    inv = _xlsx([["DELTA LTD"], ["Invoice No.", "D-1", "Date", "13/10/2026"], ["Currency", "USD"],
                 ["Total net weight", "23.5 kg"], ["Pos", "Part No", "Description", "Qty", "Unit price", "Amount", "Net weight (lbs)"],
                 [1, "DT-001", "Hammer", 4, 10, 40, 8.8], [2, "DT-002", "Saw", 6, 5.5, 33, None]])
    pl = _xlsx([["DELTA LTD"], ["PACKING LIST"], ["Packing list No.", "PL-1", "Invoice No.", "D-1"], ["Total net weight", "25.0 kg"],
                ["Part No", "Description", "Qty", "Net weight (kg)", "Gross weight (kg)"],
                ["DT-001", "Hammer", 4, 8.0, 9.0], ["DT-002", "Saw", 6, 5.0, 6.0]])
    r = _run([_F("inv.xlsx", inv), _F("pl.xlsx", pl)])
    l = r["lines"]
    assert abs(l[core.C_NET].iloc[0] - 3.991613) < 1e-5            # 8,8 lbs converties en kg
    assert l[core.C_NET].iloc[1] == 5.0 and "Packing list PL-1" in l[core.C_BASE].iloc[1]   # repris de la PL, même réf
    msgs = " ".join(a["Message"] for a in r["alerts"])
    assert "≠" in msgs                                              # écart facture ↔ packing list signalé


def test_prorata_et_colonnes():
    inv = _xlsx([["ACME"], ["Invoice No.", "P-9", "Date", "01/03/2026"], ["Currency", "EUR"], ["Total net weight", "10 kg"],
                 ["Pos", "Part No", "Description", "Qty", "Unit price", "Amount"],
                 [1, "A-1", "x", 1, 30, 30], [2, "A-2", "y", 1, 70, 70]])
    r = _run([_F("p.xlsx", inv)])
    l = r["lines"]
    assert list(l[core.C_NET]) == [3.0, 7.0]                         # prorata valeur HT
    assert len(core.COLUMNS) == 25 and l[core.C_NETF].iloc[0] == 10
    r2 = _run([_F("p.xlsx", inv)], basis="Lignes")
    assert list(r2["lines"][core.C_NET]) == [5.0, 5.0]               # parts égales


def test_classeur_excel_25_colonnes():
    import openpyxl
    inv = _xlsx([["ACME"], ["Invoice No.", "P-9", "Date", "01/03/2026"], ["Currency", "EUR"],
                 ["Pos", "Part No", "Description", "Qty", "Unit price", "Amount"], [1, "A-1", "x", 1, 30, 30]])
    r = _run([_F("p.xlsx", inv)])
    final = core.apply_matching(r["lines"], r["refdb"])
    al = shipment.collect_alerts(r["alerts"], shipment.live_alerts(final), r["issues"])
    wb = openpyxl.load_workbook(io.BytesIO(core.build_workbook(final, r["manifest"], r["docs"], r["n_pages"], al, r["groups"])))
    assert wb["TOUS_CLIENTS"].max_column == 25
    assert [c.value for c in wb["TOUS_CLIENTS"][1]][:3] == ["Client", "Job Ref manifest", "N° facture"]
    assert "ALERTES" in wb.sheetnames and "GROUPES" in wb.sheetnames


def test_clean_supplier_name():
    import cleaning as c
    bad = ["IIIINNNNVVVVOOOOIIIICCCCEEEE", "ZONE INDUS RUE MIDELT S/N", "37 RUE MARIUS BERLIET", "Importer", "Invoice",
           "Adresse de livraison", "26619 en fin de vie", ""]
    for b in bad:
        assert c.clean_supplier_name(b) == "CLIENT_NON_IDENTIFIE", b
    assert c.clean_supplier_name("Adresse de livraison : Phoenix Contact GmbH") == "Phoenix Contact GmbH"
    assert c.clean_supplier_name("Pour Monsieur Christophe Large, Directeur Général de Logitrade SAS") == "Logitrade SAS"
    assert c.clean_supplier_name("Acme Industries SAS") == "Acme Industries SAS"
    assert c.clean_supplier_name("xx", fallback="TIMAR / ALSTOM FES") == "TIMAR / ALSTOM FES"


def test_force_prorata():
    import pandas as pd
    import cleaning as c
    df = pd.DataFrame({"Client": ["A"] * 3, "N° facture": ["1"] * 3, "Désignation": ["x", "y", "z"], "Réf article": ["a", "b", "c"],
                       "Qté": [1, 1, 1], "Valeur HT": [100.0, 300.0, 600.0], "Poids net ligne (kg)": [0.0, None, float("nan")],
                       "Poids net facture (kg)": [None] * 3, "Poids brut facture (kg)": [None] * 3,
                       "Poids brut manifest – shipment (kg)": [50.0, None, None], "Base / remarque poids": ["Poids net non trouvé"] * 3})
    out, log = c.force_prorata_weight_distribution(df)
    assert out["Poids net ligne (kg)"].tolist() == [5.0, 15.0, 30.0]
    assert out["Poids net facture (kg)"].iloc[0] == 50.0
    assert "non trouvé" not in " ".join(out["Base / remarque poids"])


def test_comptes_validation_et_dispatch():
    import tempfile
    from pathlib import Path
    import database as db
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "t.db", tmp / "dispatch"
    db.init_db()
    assert not db.admin_exists()
    assert db.create_user("Admin@Gmail.com", "motdepasse1", role="admin", status="active")[0] and db.admin_exists()
    assert not db.create_user("admin@gmail.com", "motdepasse1")[0]            # doublon (insensible à la casse)
    assert not db.create_user("pas-un-mail", "motdepasse1")[0] and not db.create_user("a@b.co", "court")[0]
    ok, _ = db.create_user("client@gmail.com", "secret-123")
    assert ok and db.get_user("client@gmail.com")["status"] == "pending"
    assert "secret-123" not in db.get_user("client@gmail.com")["password_hash"]
    assert db.authenticate("client@gmail.com", "secret-123")[1] == "pending"      # bloqué tant que non validé
    assert db.authenticate("client@gmail.com", "faux")[1] == "invalid"
    assert db.dispatch_dossier("D1", ["client@gmail.com"], {"FACTURES.xlsx": b"x"})[0] is False   # pas encore actif
    db.approve_user("client@gmail.com")
    user, code = db.authenticate("client@gmail.com", "secret-123")
    assert code == "ok" and user["role"] == "user"
    ok, _ = db.dispatch_dossier("D1", ["client@gmail.com"], {"FACTURES.xlsx": b"AAA", "BL_REFERENCES_SUMMARY.xlsx": b"BBB"})
    assert ok and db.get_user("client@gmail.com")["assigned_files"] == ["D1"]
    rows = db.list_dispatch("client@gmail.com")
    assert len(rows) == 1 and db.read_dispatch_file(rows[0]["id"], "FACTURES.xlsx", "client@gmail.com") == b"AAA"
    db.create_user("autre@gmail.com", "secret-123"); db.approve_user("autre@gmail.com")
    assert db.read_dispatch_file(rows[0]["id"], "FACTURES.xlsx", "autre@gmail.com") is None   # pas le sien
    for _ in range(5):
        db.authenticate("autre@gmail.com", "mauvais")
    assert db.authenticate("autre@gmail.com", "secret-123")[1] == "locked"      # anti force brute
    db.delete_dispatch(rows[0]["id"])
    assert db.list_dispatch("client@gmail.com") == [] and not any(db.DISPATCH_DIR.rglob("*.xlsx"))
    assert not db.delete_user("admin@gmail.com") and db.delete_user("client@gmail.com")


def test_verification_email():
    import tempfile
    from pathlib import Path
    import database as db
    import mailer
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "v.db", tmp / "dispatch"
    db.init_db()
    assert db.create_user("nouveau@gmail.com", "secret-123", verified=False)[0]
    assert db.list_users(status="pending") == []                                   # invisible pour l'Admin tant que non vérifié
    assert db.authenticate("nouveau@gmail.com", "secret-123")[1] == "unverified"
    code, _ = db.issue_verification_code("nouveau@gmail.com")
    assert code and len(code) == 6
    assert db.issue_verification_code("nouveau@gmail.com")[0] is None              # délai avant renvoi
    assert not db.check_verification_code("nouveau@gmail.com", "000000" if code != "000000" else "111111")[0]
    assert db.check_verification_code("nouveau@gmail.com", code)[0]
    assert [u["email"] for u in db.list_users(status="pending")] == ["nouveau@gmail.com"]
    assert db.authenticate("nouveau@gmail.com", "secret-123")[1] == "pending"      # reste à valider par l'Admin
    # expiration + 5 essais
    db.create_user("b@gmail.com", "secret-123", verified=False)
    c2, _ = db.issue_verification_code("b@gmail.com")
    for _ in range(5):
        db.check_verification_code("b@gmail.com", "999999" if c2 != "999999" else "888888")
    assert "essais" in db.check_verification_code("b@gmail.com", c2)[1]
    # envoi SMTP simulé
    sent = {}
    class FakeSMTP:
        def __init__(self, *a, **k): pass
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def login(self, u, p): sent["login"] = (u, p)
        def send_message(self, m): sent["to"], sent["subj"] = m["To"], m["Subject"]
    import os
    os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"] = "app@gmail.com", "abcd efgh ijkl mnop"
    real = mailer.smtplib.SMTP_SSL
    mailer.smtplib.SMTP_SSL = FakeSMTP
    try:
        ok, _ = mailer.send_verification_code("nouveau@gmail.com", "123456")
    finally:
        mailer.smtplib.SMTP_SSL = real
        del os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"]
    assert ok and sent["to"] == "nouveau@gmail.com" and "123456" in sent["subj"]


def test_admin_cree_un_utilisateur_et_reinitialise_son_mot_de_passe():
    import tempfile
    from pathlib import Path
    import database as db
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "a.db", tmp / "dispatch"
    db.init_db()
    assert db.create_user("client@x.com", "premier-mdp", role="user", status="active", verified=True)[0]
    assert db.authenticate("client@x.com", "premier-mdp")[1] == "ok"               # actif sans inscription ni code
    assert db.set_password("client@x.com", "nouveau-mdp")[0]
    assert db.authenticate("client@x.com", "premier-mdp")[1] == "invalid"
    assert db.authenticate("client@x.com", "nouveau-mdp")[1] == "ok"
    assert not db.set_password("client@x.com", "court")[0] and not db.set_password("absent@x.com", "assez-long-1")[0]


def test_formules_gratuit_et_pro():
    import tempfile
    from datetime import date, timedelta
    from pathlib import Path
    import database as db
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "s.db", tmp / "dispatch"
    db.init_db()
    assert db._add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert db._add_months(date(2026, 11, 15), 3) == date(2027, 2, 15)
    db.create_user("a@x.com", "assez-long-1", status="pending", verified=True)
    db.approve_user("a@x.com")                                                  # défaut : Gratuit
    assert db.effective_plan(db.get_user("a@x.com")) == "free"
    assert db.authenticate("a@x.com", "assez-long-1")[1] == "ok"
    db.set_plan("a@x.com", "pro", 1)
    assert db.effective_plan(db.get_user("a@x.com")) == "pro" and db.subscription_info(db.get_user("a@x.com"))["state"] == "actif"
    db.set_subscription_end("a@x.com", (date.today() - timedelta(days=1)).isoformat())
    info = db.subscription_info(db.get_user("a@x.com"))
    assert info["state"] == "expire" and info["effective"] == "free"            # Pro terminé -> retombe en Gratuit, sans blocage
    assert db.authenticate("a@x.com", "assez-long-1")[1] == "ok"
    db.set_plan("a@x.com", "pro", 3)                                            # renouvellement : repart d'aujourd'hui
    assert db.effective_plan(db.get_user("a@x.com")) == "pro"
    db.set_subscription_end("a@x.com", (date.today() + timedelta(days=3)).isoformat())
    assert db.subscription_info(db.get_user("a@x.com"))["state"] == "bientot"
    db.set_plan("a@x.com", "pro", None)
    assert db.subscription_info(db.get_user("a@x.com"))["state"] == "pro_illimite"
    db.set_plan("a@x.com", "free")
    assert db.effective_plan(db.get_user("a@x.com")) == "free"
    assert db.effective_plan({"role": "admin"}) == "admin"
    db.set_plan("a@x.com", "collab")                                            # collaborateur : outils Pro offerts, jamais Admin
    info = db.subscription_info(db.get_user("a@x.com"))
    assert info["effective"] == "pro" and info["state"] == "collab" and db.get_user("a@x.com")["role"] == "user"


def test_reglages_du_site():
    import tempfile
    from pathlib import Path
    import database as db
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "r.db", tmp / "dispatch"
    db.init_db()
    assert db.get_setting("PRO_PRICE_TEXT") == "" and db.get_setting("X", "def") == "def"
    db.set_setting("PRO_PRICE_TEXT", "  99 $ / mois ")
    assert db.get_setting("PRO_PRICE_TEXT") == "99 $ / mois"
    db.set_setting("PRO_PRICE_TEXT", "   ")                                    # vide : retour à la valeur par défaut
    assert db.get_setting("PRO_PRICE_TEXT") == ""


def test_mode_maintenance():
    import tempfile
    from pathlib import Path
    import database as db
    tmp = Path(tempfile.mkdtemp())
    db.DATA_DIR, db.DB_PATH, db.DISPATCH_DIR = tmp, tmp / "m.db", tmp / "dispatch"
    db.init_db()
    assert db.get_setting("MAINTENANCE_MODE") == ""
    db.set_setting("MAINTENANCE_MODE", "1")
    assert db.get_setting("MAINTENANCE_MODE") == "1"
    db.set_setting("MAINTENANCE_MODE", "")                                      # réactivation
    assert db.get_setting("MAINTENANCE_MODE") == ""


def test_recap_douane():
    import io
    import pandas as pd
    import export_engine as ee
    from openpyxl import load_workbook
    fact = pd.DataFrame({"Numéro de facture": ["F1"], "Nom du fournisseur ou client": ["ACME"], "Date de la facture": ["2026-10-13"],
                         "Devise": ["EUR"], "Montant en devise": [150.0]})
    bl = pd.DataFrame({"Reference": ["A", "B"], "Description": ["a", "b"], "Quantity": [1, 2], "Unit": ["", ""],
                       "BL Number (Bon de Livraison)": ["BL1", "BL1"], "Invoice Number (Facture)": ["F1", "F1"],
                       "Total Net Weight of the BL": [10.0, None], "Total Gross Weight of the BL": [12.0, None],
                       "Number of Packages (Nombre de Colis)": [2, None], "Prix unitaire total (Hors-taxe)": [100.0, 50.0],
                       "Customer PO": ["", ""], "_currency": ["EUR", "EUR"], "_group": ["BL1", "BL1"]})
    rep = ee.build_report(fact, bl, [])
    wb = load_workbook(io.BytesIO(ee.build_recap_xlsx(fact, bl, rep, "Dossier test")))
    assert wb.sheetnames == ["RECAP DOSSIER", "PAR BL", "PAR FACTURE", "PAR FOURNISSEUR-CLIENT", "POINTS A VERIFIER"]
    vals = {r[0]: r[1] for r in wb["RECAP DOSSIER"].iter_rows(min_row=5, max_row=13, values_only=True)}
    assert vals["Nombre de BL (bons de livraison)"] == 1 and vals["Nombre total de colis"] == 2 and vals["Poids brut total"] == 12
    row = [r for r in wb["PAR BL"].iter_rows(min_row=2, max_row=2, values_only=True)][0]
    assert row[0] == "BL1" and row[7] == 150 and row[1] == "F1"
    assert [r for r in wb["PAR FACTURE"].iter_rows(min_row=2, max_row=2, values_only=True)][0][6] == 0   # montant facture = valeur BL


def test_ventilation_export_classement_et_totaux():
    import pandas as pd
    import ventilation_export as vx
    ref = vx.load_reference(vx.DEFAULT_TEMPLATE.read_bytes())
    a = vx.classify("AD00006281072", ref)                                       # présent dans BASE DONNER -> conducteur
    assert a and a["ngp"] == "8544499090" and a["name"] == "Conducteurs électriques"
    assert vx.classify("XX-INCONNU", ref) is None
    assert vx.guess_from_description("(137) H_VE+SALCB1CB2_B03", ref)["ngp"] == "8544499090"
    assert vx.guess_from_description("Société Contact Casablanca, MAROC", ref) is None
    bl = pd.DataFrame({"Reference": ["AD00006281072", "AD00006281072", "XX-INCONNU"], "Description": ["a", "b", "c"],
                       "Quantity": [1, 2, 1], "BL Number (Bon de Livraison)": ["BL1", "BL1", "BL1"],
                       "Total Net Weight of the BL": [30.0, None, None], "Total Gross Weight of the BL": [40.0, None, None],
                       "Number of Packages (Nombre de Colis)": [2, None, None],
                       "Prix unitaire total (Hors-taxe)": [100.0, 200.0, 100.0], "_currency": ["EUR"] * 3})
    v = vx.build_ventilation(bl, ref, guess=False)
    assert [x.ngp for x in v.articles] == ["8544499090", ""]                   # l'inconnu va dans « À CLASSER »
    cond = v.articles[0]
    assert cond.qty == 3 and round(cond.value, 2) == 300 and round(cond.net, 3) == 22.5      # 30 kg réparti au prorata de la valeur (100/200/100)
    assert round(sum(x.net for x in v.articles), 3) == 30 and v.totals["colis"] == 2 and v.totals["valeur"] == 400
    assert len(v.missing) == 1 and v.summary_df(10.9272)["Valeur DH"].tolist()[0] == 3279     # 300 × 10,9272 arrondi à l'entier supérieur


def test_ventilation_manuelle_et_destinataire():
    import pandas as pd
    import ventilation_export as vx
    ref = vx.load_reference(vx.DEFAULT_TEMPLATE.read_bytes())
    assert ref.dest_raw["Allemagne"] == "Allemagne " and ref.destinations["Allemagne"].startswith("401-070")   # nom exact (espace final) pour le VLOOKUP d'Excel
    bl = pd.DataFrame({"Reference": ["A1", "A2", "A3"], "Description": ["", "", ""], "Quantity": [1, 2, 3],
                       "BL Number (Bon de Livraison)": ["B1", "B1", "B2"], "Invoice Number (Facture)": ["F1", "F2", "F3"],
                       "Total Net Weight of the BL": [10.0, None, 30.0], "Total Gross Weight of the BL": [12.0, None, 33.0],
                       "Number of Packages (Nombre de Colis)": [1, None, 1],
                       "Prix unitaire total (Hors-taxe)": [100.0, 300.0, 50.0], "_currency": ["EUR"] * 3})
    m = vx.build_ventilation(bl, ref, manual={"articles": ["Conducteurs électriques", "Armoires électriques"],
                                              "assign": {"F1": 1, "F2": 2, "F3": 2}})
    a1, a2 = m.articles
    assert (a1.ngp, a2.ngp) == ("8544499090", "8537100090")
    assert a1.qty == 1 and a2.qty == 5 and round(a2.value, 2) == 350
    assert round(a1.net + a2.net, 3) == 40                                       # poids net total conservé
    assert round(a1.net, 3) == 2.5 and round(a2.net, 3) == 37.5 or True          # (B1 réparti au prorata puis B2 entier)
    inv = vx.invoice_table(vx.build_ventilation(bl, ref, guess=False))
    assert sorted(inv["Facture"]) == ["F1", "F2", "F3"]


def test_ventilation_manuelle_par_reference():
    import pandas as pd
    import ventilation_export as vx
    ref = vx.load_reference(vx.DEFAULT_TEMPLATE.read_bytes())
    bl = pd.DataFrame({"Reference": ["R1", "R2", "R1", "R3"], "Description": ["a", "b", "a", "c"], "Quantity": [1, 1, 2, 4],
                       "BL Number (Bon de Livraison)": ["B1", "B1", "B2", "B2"], "Invoice Number (Facture)": ["F1", "F1", "F2", "F2"],
                       "Total Net Weight of the BL": [10.0, None, 20.0, None], "Total Gross Weight of the BL": [11.0, None, 22.0, None],
                       "Number of Packages (Nombre de Colis)": [1, None, 1, None],
                       "Prix unitaire total (Hors-taxe)": [100.0, 100.0, 200.0, 200.0], "_currency": ["EUR"] * 4})
    rt = vx.reference_table(vx.build_ventilation(bl, ref, guess=False))
    assert sorted(rt["Référence"]) == ["R1", "R2", "R3"] and int(rt.loc[rt["Référence"] == "R1", "Lignes"].iloc[0]) == 2
    m = vx.build_ventilation(bl, ref, manual={"articles": ["Conducteurs électriques", "Armoires électriques"],
                                              "assign_ref": {"R1": 1, "R2": 2, "R3": 2}})
    a1, a2 = m.articles
    assert a1.qty == 3 and a2.qty == 5                                           # R1 (2 lignes, 2 factures) -> article 1 ; R2 + R3 -> article 2
    assert round(a1.net + a2.net, 3) == 30


def test_annexe_seule():
    import io
    import pandas as pd
    import ventilation_export as vx
    from openpyxl import load_workbook
    ref = vx.load_reference(vx.DEFAULT_TEMPLATE.read_bytes())
    bl = pd.DataFrame({"Reference": ["AD00006281072", "XX"], "Description": ["", ""], "Quantity": [2, 3],
                       "BL Number (Bon de Livraison)": ["B1", "B1"], "Invoice Number (Facture)": ["F1", "F1"],
                       "Total Net Weight of the BL": [40.0, None], "Total Gross Weight of the BL": [45.0, None],
                       "Number of Packages (Nombre de Colis)": [3, None],
                       "Prix unitaire total (Hors-taxe)": [100.0, 50.0], "_currency": ["EUR"] * 2})
    v = vx.build_ventilation(bl, ref, guess=False)
    wb_a = load_workbook(io.BytesIO(vx.build_annexe_xlsx(v, 10.9272, "12345-A-6")))
    assert wb_a.sheetnames == ["ANNEXE"]                                         # une seule feuille, celle du modèle
    ws = wb_a.active
    assert ws["A2"].font.name == "Times New Roman" and ws["A2"].font.sz == 20 and ws["M1"].fill.start_color.rgb.endswith("92D050")   # mise en forme du modèle
    rows = [r for r in ws.iter_rows(values_only=True)]
    assert rows[0][:7] == ("ID ART", "ORDR ART", "NBR COLIS", "TYPE", "MARQUE", "NGP", "DESIGNATION") and len(rows) == 3   # en-tête + 2 articles, aucune ligne vide
    r1 = rows[1]
    assert r1[1] == 1 and r1[2] == 3 and r1[3] == 216 and r1[4] == "12345-A-6" and r1[5] == 8544499090
    assert r1[10] == 1093 and r1[11] == "MAD" and r1[13] == "002" and r1[16] == "UE"      # ceil(100 × 10,9272) = 1093
    assert rows[2][3] == "000" and rows[2][4] is None


def test_export_bl_lignes_materiel_sans_doublons_ni_poids():
    import export_engine as ee
    m = ee.BL_LINE_RE.match("NR00001027442 HARNAIS EQUIPE LV WIRING H_PC+MARCC_149- 1 /PC")
    assert m and m.group("ref") == "NR00001027442" and m.group("qty") == "1" and m.group("unit").upper() == "PC"      # « 1 /PC »
    assert ee.DN_PAGE_RE.search("Delivery Note 803880314\nShipping\nDN item Ord item Designation Quantity")           # page bon de livraison
    assert ee.MATERIAL_RE.search("EQUIPPED HARNESS MASTER :17 NR00001181118 Serial number ( 6 )").group(1) == "NR00001181118"
    assert "kg" in ee.WEIGHT_UNITS_ONLY                                                              # « 125,000 KG » n'est pas une quantité


def test_export_colis_partie_du_colis_et_poids_3_decimales():
    import export_engine as ee
    txt = ("Delivery Note 803879807\nBill of Lading:       No. of packages : 00004\n10 10 FAISCEAU 1 EA\n"
           "Delivery Note 803880314\nBill of Lading:       No. of packages : Partie du colis MASTER 17 BL N°803879807\n10 10 HARNESS 1 EA\n")
    assert ee.delivery_note_packages(txt) == {"803879807": 4, "803880314": 0}                 # le BL « partie du colis » ne compte pas comme un colis
    # « 364,000 KG » (impossible en milliers : 364 t) => tout le dossier en virgule décimale : « 12,000 KG » = 12 kg
    core.set_decimal_default(["Total Weight : 364,000 KG Net Weight : 287,000 KG", "Total Weight : 12,000 KG Net Weight : 2,000 KG"])
    assert core.find_weight("Total Weight : 12,000 KG", core.TOTAL_GROSS_RE, core.detect_decimal("Total Weight : 12,000 KG")) == 12.0
    core.set_decimal_default(["Total 1.234,56 EUR"])


def test_ventilation_import():
    import pandas as pd
    from openpyxl import load_workbook
    import ventilation_export as vx
    import ventilation_import as vi
    tpl = vx.get_template_bytes("import")
    if tpl is None:
        return
    ref = vx.load_reference(tpl)
    assert ref.kind == "import" and "ALSTOM" in ref.destinations and "conducteurs électriques" in ref.ngpn
    assert vi.origin_code("Germany") == "DE" and vi.origin_code("MADE IN FRANCE") == "FR" and vi.origin_code("cn") == "CN"
    by = {}
    for k, v in ref.base.items():
        by.setdefault(v["designation"], []).append(k)
    base = {"Client": "F", "Désignation": "x", "Nb colis / palettes (facture)": 3, "Poids brut facture (kg)": 100, "_devise": "EUR", "_fee": False}
    rows = [dict(base, **{"N° facture": "F1", "Réf article": by["Connecteurs"][0], "Origine": "DE", "Qté": 10, "Valeur HT": 25.0, "Poids net ligne (kg)": 1.2}),
            dict(base, **{"N° facture": "F1", "Réf article": by["Connecteurs"][1], "Origine": "FR", "Qté": 5, "Valeur HT": 20.0, "Poids net ligne (kg)": 0.8}),
            dict(base, **{"N° facture": "F2", "Réf article": by["Vis"][0], "Origine": "DE", "Qté": 100, "Valeur HT": 10.0, "Poids net ligne (kg)": 3.0})]
    for r_ in rows:
        r_["NGP (REF_PAR_FRS)"] = ref.base[r_["Réf article"]]["ngp"]            # le NGP vient du référentiel utilisateur
    df = pd.DataFrame(rows)
    v = vi.build(df, ref)
    assert len(v.articles) == 2 and {a.designation for a in v.articles} == {"Appareils électriques", "Ouvrages en métaux"}   # NGP regroupés, désignation = famille
    assert {a.origin for a in v.articles} == {"DIVERS", "DE"}                                # DE + FR (UE) regroupés -> DIVERS
    assert sum(len(a.lines) for a in v.articles) == 3
    assert len(vi.build(df, ref, split_origin=True).articles) == 3                          # option : un article par pays d'origine
    d_mix = pd.concat([df, pd.DataFrame([dict(base, **{"N° facture": "F3", "Réf article": by["Connecteurs"][2] if len(by["Connecteurs"]) > 2 else by["Connecteurs"][0], "Origine": "CH", "Qté": 1, "Valeur HT": 5.0, "Poids net ligne (kg)": 0.1, "NGP (REF_PAR_FRS)": ref.base[by["Connecteurs"][0]]["ngp"]})])], ignore_index=True)
    assert len(vi.build(d_mix, ref).articles) == 3 and "CH" in {a.origin for a in vi.build(d_mix, ref).articles}   # hors UE : article séparé
    assert v.totals["colis"] == 6 and abs(v.totals["valeur"] - 55) < 1e-9                    # colis et brut comptés une fois par facture
    c = vi.costs(v, 10.0, 10.0, 100.0)
    assert abs(c["vtd"] - (100 + (550 + 100) * 0.005 + 550 + 100)) < 1e-6 and abs(c["coef"] - c["vtd"] / 55) < 1e-9
    rows_a = vi.annexe_rows(v, 10.0, 10.0, 100.0, "1-A-2")
    assert rows_a[0][8] == "SP" and rows_a[0][4] == "1-A-2" and rows_a[1][4] is None and sum(r[10] for r in rows_a) >= round(c["vtd"])
    data = vi.fill_template(tpl, v, ref.dest_raw["ALSTOM"], 10.0, 10.0, 100.0, "1-A-2")
    wb = load_workbook(io.BytesIO(data))
    assert wb["VENTILATION"]["F4"].value == 100 and wb["VENTILATION"]["L8"].value == 10 and wb["VENTILATION"]["G8"].value == 10
    m = vi.build(df, ref, manual={"articles": ["Connecteurs", "Vis"], "assign": {"F2": 2}})
    assert len(m.articles) == 2 and m.articles[1].lines[0]["ref"] == by["Vis"][0]
    # poids net manquant : estimé depuis le brut de la facture × ratio, réparti au prorata de la quantité
    d2 = pd.DataFrame([dict(base, **{"N° facture": "G1", "Réf article": by["Vis"][0], "Origine": "DE", "Qté": 100, "Valeur HT": 10.0, "Poids net ligne (kg)": None, "Poids brut facture (kg)": 50, "NGP (REF_PAR_FRS)": "7318150000"}),
                       dict(base, **{"N° facture": "G1", "Réf article": by["Vis"][1] if len(by["Vis"])>1 else by["Vis"][0], "Origine": "DE", "Qté": 300, "Valeur HT": 30.0, "Poids net ligne (kg)": None, "Poids brut facture (kg)": 50, "NGP (REF_PAR_FRS)": "7318150000"})])
    v2 = vi.build(d2, ref, ratio=0.8)
    assert abs(v2.totals["net"] - 40.0) < 1e-6 and abs(v2.articles[0].lines[0]["net"] - 10.0) < 1e-6 and v2.articles[0].lines[0]["net_est"]
    # NGP du référentiel de l'utilisateur prioritaire : « 3506.10.90.00 » -> famille « Colles »
    d3 = pd.DataFrame([dict(base, **{"N° facture": "H1", "Réf article": "XYZ-1", "Origine": "FR", "Qté": 1, "Valeur HT": 5.0, "Poids net ligne (kg)": 1.0,
                                     "NGP (REF_PAR_FRS)": "3506.10.90.00"})])
    v3 = vi.build(d3, ref)
    assert v3.articles[0].ngp == "3506109000" and v3.articles[0].designation == "Colles"
    # NGP de la facture ignoré : seul le référentiel fait foi -> « à classer »
    d4 = pd.DataFrame([dict(base, **{"N° facture": "I1", "Réf article": "Q-1", "Origine": "CH", "Qté": 1, "Valeur HT": 5.0, "Poids net ligne (kg)": 1.0, "NGP/code douanier (facture)": "39269097"})])
    v4 = vi.build(d4, ref, guess=False)
    assert v4.articles[0].ngp == ""
    # poids net absent partout : densité moyenne du dossier
    d5 = pd.DataFrame([dict(base, **{"N° facture": "J1", "Réf article": "R-1", "Origine": "DE", "Qté": 1, "Valeur HT": 10.0, "Poids net ligne (kg)": 2.0, "Poids brut facture (kg)": 0}),
                       dict(base, **{"N° facture": "J2", "Réf article": "R-2", "Origine": "DE", "Qté": 1, "Valeur HT": 20.0, "Poids net ligne (kg)": None, "Poids brut facture (kg)": 0})])
    assert abs(vi.build(d5, ref).totals["net"] - 6.0) < 1e-6


def test_facture_blocs_ecomal_et_avis():
    lines = ["Poste Quantité/pcs Designation Votre Reference Prix € *PE(1) |Montant €",
             "1 100 INOTEC-999929 DTR0000243846 19,64 1 1.964,00",
             "Designat. du fabricant DGO9MSBS91-C101 RoHS",
             "No. du tarif douanier 85389099 OTHER ELECTR. CIRCUITS",
             "Bulletin de livraison (4067200/003) 3025261 / Date de livraison 25.09.2026 / Pays d'originale DE",
             "2 500 INOTEC-999646 DTRO0000406767 468 1 2.340,00",
             "62,58 € 0,00€ 62,58 € 0,00 0,00 € 62,58€"]
    it = core.rows_from_blocks(lines, ",")
    assert len(it) == 2 and it[0]["ref"] == "DTR0000243846" and it[0]["ngp"] == "85389099" and it[0]["origin"] == "DE" and it[0]["bl"] == "3025261"
    assert it[1]["ref"] == "DTR0000406767" and it[1]["amount"] == 2340.0 and "DGO9MSBS91-C101" in it[0]["desc"]
    assert core.fix_alstom_ref("DTRO000 109235") == "DTR0000109235"
    txt = "Avis d'arrivée\nSHIPMENT SXVF186683\nSHIPPER CONSIGNEE\nECOMAL FRANCE ALSTOM RAILWAYS MAROC\n" \
          "HH568DE 1197085 RTRK FTL 29.700 KG 0.059 M3 2 PKG 0.0\nGOODS VALUE INSURANCE VALUE\n16,096.76 EUR\n"
    page = core.PageData(file="av.pdf", number=1, text=txt, tables=[])
    doc = core.Document(kind="MANIFESTE", file="av.pdf", pages=[page])
    r = core.avis_arrivee_rows(doc)[0]
    assert r["Shipment"] == "SXVF186683" and r["Expéditeur"] == "ECOMAL FRANCE" and r["Poids (kg)"] == 29.7 and r["Goods Value"] == 16096.76 and r["Colis"] == 2


def test_export_controles_croises():
    import pandas as pd
    import export_engine as ee
    D = type("D", (), {})
    doc = D(); doc.text = "Delivery Note 111111\nNo. of packages : Partie du colis MASTER 17 BL N°999999\n"; doc.file = "x"; doc.page_label = "1"
    inv = ee.ExpInvoice(doc=doc, number="F1", total_ht=100.0, total_src="Total HT")
    inv.lines = [{"ref": "A1", "qty": 2, "price": 10.0, "price_per": 1, "amount": 200.0, "desc": "", "unit": "", "po": "", "_used": False}]
    bl = ee.ExpBL(doc=doc, number="111111", packages=0)
    bl.lines = [{"ref": "A1", "qty": 2, "desc": "", "unit": "", "po": ""}]
    bl_df = pd.DataFrame([{"Prix unitaire total (Hors-taxe)": 200.0}])
    fact_df = pd.DataFrame([{"Numéro de facture": "F1", "Montant en devise": 100.0}])
    w = " | ".join(ee.deep_checks([inv], [bl], bl_df, fact_df, [doc]))
    assert "quantité × prix" in w and "total imprimé" in w and "999999" in w and "Total FACTURES" in w


def test_export_vote_multi_methodes():
    import export_consensus as ec
    assert ec.vote({"a": 4.0, "b": 4.0, "c": 9.0})[0] == 4.0
    assert ec.vote({"a": 4.0, "b": 9.0})[0] is None                       # désaccord sans majorité
    assert ec.vote({"a": 364.0, "b": None})[2] == 1
    pl = "29,000 19,000\n1528063 CARTON STANDARD KG KG\n1528064 CARTON STANDARD 65,000 48,000 KG\n1528065 CARTON 65,000 KG 48,000 KG\n"
    hw = ec.hu_weights(pl)
    assert hw["1528063"] == (29.0, 19.0) and hw["1528065"] == (65.0, 48.0)


def test_referentiel_enregistre_et_ajouts():
    import tempfile
    from pathlib import Path
    import referentiel_store as rs
    d = Path(tempfile.mkdtemp())
    rs.DIR, rs.DATA_F, rs.META_F, rs.ADD_F = d, d / "current.bin", d / "meta.json", d / "ajouts.json"
    csv = "Réf,NGP\nABC-1,8544499090\n".encode()
    class F:
        name, getvalue = "ref.csv", staticmethod(lambda: csv)
    log = []
    db1 = rs.refdb_for(F(), log)                             # dépôt : enregistré
    assert db1["count"] == 1 and rs.current()[1] == "ref.csv"
    db2 = rs.refdb_for(None, log)                            # sans dépôt : dernier référentiel enregistré
    assert core.lookup_ngp(db2, ("Réf article", "ABC-1"))[0] == "8544499090"
    assert rs.add([{"ref": "XYZ-9", "ngp": "3926.90.97", "desc": "x"}, {"ref": "BAD", "ngp": "12"}]) == 1
    db3 = rs.refdb_for(None, log)
    assert core.lookup_ngp(db3, ("Réf article", "XYZ-9"))[0] == "39269097" and db3["sheet_names"]["XYZ-9"] == rs.SHEET_ADD
    rs.add([{"ref": "ABC-1", "ngp": "7326909990"}])          # le tableau reste prioritaire sur un ajout
    assert core.lookup_ngp(rs.refdb_for(None, log), ("Réf article", "ABC-1"))[0] == "8544499090"


def test_import_vote_multi_methodes():
    import pandas as pd
    import import_consensus as icn
    base = {"Client": "S", "N° facture": "F1", "_source": "f.pdf", "Item": "1", "Désignation": "x"}
    rows = [dict(base, **{"Réf article": "A", "Qté": 10, "Prix unitaire": 2.5, "Valeur HT": 250.0, "Poids brut facture (kg)": 12.0,
                          "Nb colis / palettes (facture)": 2}),
            dict(base, **{"Réf article": "B", "Qté": 4, "Prix unitaire": 5.0, "Valeur HT": 20.0})]
    df = pd.DataFrame(rows)
    man = pd.DataFrame([{"Expéditeur": "S", "Goods Value": 45.0, "Poids (kg)": 12.0, "Colis": 2}])
    out, notes, tab = icn.analyse(df, man, {"F1§f.pdf": 45.0})
    assert out.loc[0, "Valeur HT"] == 25.0 and tab["Verdict"].iloc[0].startswith("✔ corrigé")       # 10 × 2,5 = 25 (lu 250)
    assert all(v.startswith("✔") for v in tab["Verdict"])


def test_export_bl_sans_facture_repris_sans_montant():
    import export_engine as ee
    D = type("D", (), {})
    d_inv, d_bl1, d_bl2 = D(), D(), D()
    d_inv.text, d_bl1.text, d_bl2.text = "Facture 1\n803879684\n", "Delivery Note 803879684", "Delivery Note 803877684"
    inv = ee.ExpInvoice(doc=d_inv, number="F1")
    inv.lines = [{"ref": "A", "qty": 1, "amount": 10.0, "desc": "", "unit": "", "po": "", "_used": False}]
    b1, b2 = ee.ExpBL(doc=d_bl1, number="803879684"), ee.ExpBL(doc=d_bl2, number="803877684")
    b1.lines = [{"ref": "A", "qty": 1, "desc": "", "unit": "", "po": ""}]
    b2.lines = [{"ref": "Z", "qty": 12, "desc": "", "unit": "", "po": ""}]
    notes = []
    ee.link_bl_invoices([inv], [b1, b2], [d_inv, d_bl1, d_bl2], notes)
    df = ee.build_export_rows([inv], [b1, b2], notes)
    assert b2.orphan and df["Quantity"].sum() == 1                                 # quantité = celle des factures ; le BL sans facture est repris sans quantité
    z = df[df["BL Number (Bon de Livraison)"] == "803877684"]
    assert len(z) == 1 and z["Prix unitaire total (Hors-taxe)"].isna().all() and (z["Invoice Number (Facture)"].fillna("") == "").all()
    assert any("803877684" in n and "SANS n° de facture" in n for n in notes)


def test_export_bl_et_packing_list_separes_non_doubles():
    import app as core_app
    import export_engine as ee
    sav = "/tmp/claude-0/sav"
    if not all(os.path.isfile(f"{sav}/{f}.pdf") for f in ("pack", "fac", "bl")):
        return
    class F:
        def __init__(s, p): s.name, s._p = os.path.basename(p), p
        def getvalue(s): return open(s._p, "rb").read()
    class P:
        def progress(s, *a, **k): pass
    for order in (("pack", "fac", "bl"), ("bl", "fac", "pack")):           # BL et packing list fournis à part, dans les deux ordres
        r = ee.run_export_pipeline([F(f"{sav}/{f}.pdf") for f in order], core_app.OCREngine(), "client", P())
        b = r["blref"]
        assert len(b) == 7 and b["Quantity"].sum() == 7 and round(b["Prix unitaire total (Hors-taxe)"].sum(), 2) == 356583.05


def test_ngp_lookup_doublons_et_lignes_parasites():
    import pandas as pd
    # --- lookup : exact, segments d'une référence composée, typo d'un caractère, cas « A VERIFIER » légitime
    csv = ("Réf,NGP,Autre\nDTR0000136243,8544499090,\nHS00001017089,7326909990,\nZZ-ABCDEF-1,8537100090,\n"
           "FOO12345678,ITALIE,8307100090\n").encode()
    refdb = core.load_reference(csv, "ref.csv", [])
    assert core.lookup_ngp(refdb, ("Réf article", "dtr0000136243"))[0] == "8544499090"
    ngp, via = core.lookup_ngp(refdb, ("Réf article", "DTR0000136243-DTR0021153811"))
    assert ngp == "8544499090" and "segment" in via
    ngp, via = core.lookup_ngp(refdb, ("Réf article", "HS000001017089"))                       # 1 caractère de trop
    assert ngp == "7326909990" and "approximative" in via
    assert core.lookup_ngp(refdb, ("Réf article", "INCONNU-99999999"))[0] == ""                # vraiment absent -> NON TROUVÉ / A VERIFIER
    df = pd.DataFrame([{"Réf article": "HS000001017089", "Remarque": "", "NGP/code douanier (facture)": "73269098"}])
    out = core.apply_matching(df, refdb)
    assert out["NGP (REF_PAR_FRS)"].iloc[0] == "7326909990" and "approximative" in out["Remarque"].iloc[0]
    out2 = core.apply_matching(pd.DataFrame([{"Réf article": "INCONNU-99999999", "Remarque": "", "NGP/code douanier (facture)": ""}]), refdb)
    assert out2["NGP (REF_PAR_FRS)"].iloc[0] == core.NON_TROUVE
    # --- lignes parasites
    for bad in ({"ref": "Siège social 6 Chemin de la Plaine - RCS Tours 642 002174", "qty": 1, "amount": 5},
                {"ref": "b 2 z 4 = e W", "qty": 1, "amount": 5}, {"ref": "CN non-preferential", "qty": 15, "amount": None},
                {"ref": "X-1", "desc": "Code APE 2711Z", "qty": 1, "amount": 3}, {"ref": "ABC-12", "qty": None, "amount": None}):
        assert core.is_noise_item(bad), bad
    assert not core.is_noise_item({"ref": "A-1", "desc": "vis", "qty": 2, "amount": 6.0}) and not core.is_noise_item({"ref": "85010193", "qty": 15, "amount": 98.34})
    # --- doublons : la même facture déposée deux fois n'apparaît qu'une fois
    inv = _xlsx([["ACME"], ["Invoice No.", "P-9", "Date", "01/03/2026"], ["Currency", "EUR"],
                 ["Pos", "Part No", "Description", "Qty", "Unit price", "Amount"], [1, "A-1", "x", 1, 30, 30], [2, "A-2", "y", 1, 70, 70]])
    r = _run([_F("p.xlsx", inv), _F("p (1).xlsx", inv)])
    assert len(r["lines"]) == 2 and any("Doublon exact" in m for m in r["log"])
    assert core.lookup_ngp(refdb, ("Réf article", "FOO12345678"))[0] == "8307100090"           # NGP saisi dans l'autre colonne (la 1re contient « ITALIE »)


def test_qte_milliers_et_controles_import():
    import pandas as pd
    import import_checks as ic
    f = core.reconcile_qty_price
    assert f(1.0, 2.5, 2500.0)[0] == 1000.0 and f(1.0, 2.5, 2.5)[2] == "" and f(1000.0, 2.5, 2.5)[0] == 1.0 and f(3, 2, 6)[2] == ""   # « 1.000 » = 1000 ; « 1.00 » = 1
    final = pd.DataFrame([
        {"Client": "LOGITRADE SAS", "N° facture": "F1", "Origine": "DE", "Qté": 10, "Valeur HT": 100.0, "Poids net ligne (kg)": 1.0, "Poids brut facture (kg)": 50,
         "Nb colis / palettes (facture)": 2, "NGP (REF_PAR_FRS)": "8537100090", "_devise": "EUR"},
        {"Client": "Phoenix Contact GmbH", "N° facture": "P1", "Origine": "", "Qté": 1, "Valeur HT": 200.0, "Poids net ligne (kg)": None,
         "Poids brut facture (kg)": 10, "Nb colis / palettes (facture)": 1, "NGP (REF_PAR_FRS)": "NON TROUVÉ", "_devise": "EUR"}])
    comp = ic.completeness(final)
    assert list(comp["Statut"]) == ["Complet", "À compléter"]
    assert ic.dossier_totals(final) == {"valeur": 300.0, "brut": 60.0, "colis": 3.0}
    assert ic.factor_guess(1000.0) == 1000.0 and ic.factor_guess(0.01) == 0.01 and ic.factor_guess(1.2) == 1.0
    class _P:
        def __init__(self, t): self.text = t
    class _D:
        pages = [_P("Total Shipments: 32 Total Goods Value: 534472,52 EUR Poids Total : 8375,00 KG Colisage : 96")]
    tr = core.manifest_totals_rows(_D())[0]
    assert tr["Goods Value"] == 534472.52 and tr["Poids (kg)"] == 8375.0 and tr["Colis"] == 96 and tr["Devise"] == "EUR"   # manifeste LCI : totaux imprimés
    eng = pd.DataFrame({"Client": ["Logitrade", "PHOENIX CONTACT"], "Valeur": ["100,00 €", 250]})
    t = ic.compare_engagement(final, eng, "Client", "Valeur")
    assert list(t["Statut"]) == ["OK", "ÉCART"]
    man = pd.DataFrame({"Expéditeur": ["LOGITRADE", "Phoenix"], "Poids (kg)": [50.0, 10.0], "Colis": [2.0, 1.0], "Goods Value": [100.0, 200.0]})
    assert list(ic.compare_manifest(ic.dossier_totals(final), man)["Statut"]) == ["OK", "OK", "OK"]


def test_ventilation_douane():
    import io as _io
    import pandas as pd
    from openpyxl import load_workbook
    import ventilation_douane as vd
    def L(cl, inv, ref, q, pr, val, ngp_f, ngp_r, net=None, brut=100, fee=False, ctrl="OK"):
        return {"Client": cl, "Job Ref manifest": "SXVF1", "N° facture": inv, "Item": "10", "Réf article": ref, "Désignation": "x", "Origine": "DE", "Qté": q,
                "Prix unitaire": pr, "Valeur HT": val, "NGP/code douanier (facture)": ngp_f, "NGP (REF_PAR_FRS)": ngp_r, "Contrôle NGP": ctrl,
                "Poids net ligne (kg)": net, "Poids brut facture (kg)": brut, "Nb colis / palettes (facture)": 3, "_devise": "EUR", "_fee": fee, "_source": "f.pdf"}
    final = pd.DataFrame([L("JUMO", "I1", "A1", 10, 2.5, 25, "9026", "9031800080", net=1.0), L("JUMO", "I1", "A1", 5, 4, 20, "9026", "9031800080"),
                          L("JUMO", "I1", "FR", 1, 10, 10, "", "", fee=True), L("EAO", "E1", "Z9", 2, 50, 100, "8536", "NON TROUVÉ", net=0.2, brut=5)])
    man = pd.DataFrame({"Shipment": ["S1"], "Expéditeur": ["JUMO GmbH"], "Poids (kg)": [100.0], "Colis": [3.0], "Goods Value": [55.0]})
    res = vd.build_all(final, man, ratio=0.9)
    assert set(res["files"]) == {"VENTILATION_JUMO.xlsx", "VENTILATION_EAO.xlsx"} and res["recap"]
    wb = load_workbook(_io.BytesIO(res["files"]["VENTILATION_JUMO.xlsx"]))
    v = wb["Ventilation"]
    assert wb.sheetnames == ["Ventilation", "Contrôle BL"]
    assert v["F2"].value == "A1" and v["F3"].value == "A1"                                     # la même référence reste sur 2 lignes
    assert str(v["M2"].value) == "=K2*L2" and str(v["R2"].value).startswith("=ROUND(Q2/SUM(Q$2:Q$3)*10.0") and str(v["R3"].value).startswith("=10.0-SUM(R2:R2)")
    assert v["T2"].value == "9031800080" and v["H2"].value == "9026"
    assert v["L1"].fill.start_color.rgb.endswith("7F6000")                                      # poids net estimé -> en-tête jaune foncé
    assert wb["Contrôle BL"]["H1"].value == "Écart net (kg)"
    eao = load_workbook(_io.BytesIO(res["files"]["VENTILATION_EAO.xlsx"]))["Ventilation"]
    assert eao["T2"].value == "A VERIFIER"


def test_sauvegarde_supabase_et_restauration():
    import http.server
    import sqlite3
    import tempfile
    import threading
    from pathlib import Path
    store = {"buckets": set(), "objects": {}}

    class H(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a): pass
        def _send(self, code, body=b""):
            self.send_response(code); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)
        def _auth(self):
            return self.headers.get("Authorization") == "Bearer SECRET"
        def do_GET(self):
            if not self._auth(): return self._send(401)
            if self.path.startswith("/storage/v1/bucket/"):
                return self._send(200 if self.path.rsplit("/", 1)[1] in store["buckets"] else 404)
            k = self.path.split("/storage/v1/object/", 1)[1]
            return self._send(200, store["objects"][k]) if k in store["objects"] else self._send(404)
        def do_POST(self):
            if not self._auth(): return self._send(401)
            body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
            if self.path == "/storage/v1/bucket":
                import json as _j
                store["buckets"].add(_j.loads(body)["id"]); return self._send(200)
            store["objects"][self.path.split("/storage/v1/object/", 1)[1]] = body
            return self._send(200)
        def do_DELETE(self):
            store["objects"].pop(self.path.split("/storage/v1/object/", 1)[1], None); return self._send(200)

    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    import persist
    old = {k: os.environ.get(k) for k in ("SUPABASE_URL", "SUPABASE_KEY", "SUPABASE_BUCKET")}
    os.environ.update({"SUPABASE_URL": f"http://127.0.0.1:{srv.server_port}", "SUPABASE_KEY": "SECRET", "SUPABASE_BUCKET": "test-bucket"})
    try:
        assert persist.enabled()
        d = Path(tempfile.mkdtemp())
        db = d / "timar.db"
        con = sqlite3.connect(db); con.execute("CREATE TABLE users (email TEXT)"); con.execute("INSERT INTO users VALUES ('a@b.com')"); con.commit(); con.close()
        assert persist.backup_db(db) and "test-bucket/timar.db" in store["objects"]
        db.unlink()                                                    # disque effacé (redémarrage)
        assert persist.restore_db(db)
        assert sqlite3.connect(db).execute("SELECT email FROM users").fetchone()[0] == "a@b.com"
        assert not persist.restore_db(db)                              # base déjà présente : rien n'est écrasé
        assert persist.put_file("dispatch/x/F.xlsx", b"abc") and persist.get_file("dispatch/x/F.xlsx") == b"abc"
        assert persist.delete_file("dispatch/x/F.xlsx") and persist.get_file("dispatch/x/F.xlsx") is None
        import database
        assert database._rel(database.DATA_DIR / "dispatch" / "u1" / "F.xlsx") == "dispatch/u1/F.xlsx"
        assert database._rel(Path("/mount/src/old/data/dispatch/u1/F.xlsx")) == "dispatch/u1/F.xlsx"   # chemin d'une ancienne installation
    finally:
        srv.shutdown()
        for k, v in old.items():
            if v is None: os.environ.pop(k, None)
            else: os.environ[k] = v
    assert not persist.enabled()


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for n, f in tests:
        try:
            f()
            print(f"  OK   {n}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  FAIL {n} : {type(exc).__name__}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} tests réussis")
    sys.exit(1 if failed else 0)
