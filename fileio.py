# -*- coding: utf-8 -*-
"""Lecture multi-formats : Excel (.xlsx/.xlsm/.xls), CSV, Word (.docx), texte et validation des fichiers téléversés.

Chaque fichier lisible est converti en « pages » (texte + tableaux) exactement comme un PDF, de sorte que le moteur
d'extraction (classement, factures, packing lists, manifestes) fonctionne à l'identique quel que soit le format.
Toute erreur de lecture est levée sous forme de FileReadError avec un message clair destiné à l'utilisateur.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from xml.etree import ElementTree as ET

KINDS = {
    ".pdf": "pdf", ".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".csv": "csv", ".tsv": "csv",
    ".docx": "word", ".txt": "text",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".tif": "image", ".tiff": "image", ".bmp": "image",
    ".webp": "image",
}
UPLOAD_TYPES = [ext.lstrip(".") for ext in KINDS]
KIND_LABEL = {"pdf": "PDF", "excel": "Excel", "csv": "CSV", "word": "Word", "text": "Texte", "image": "Image"}
MAX_MB = 250
MAX_ROWS = 30000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class FileReadError(Exception):
    """Fichier illisible : le message est affichable tel quel à l'utilisateur."""


@dataclass
class SheetPage:
    label: str
    text: str
    tables: list = field(default_factory=list)
    decimal: str = ""          # '.' si les nombres sont de vrais nombres Excel, '' = à détecter
    n_rows: int = 0


def file_kind(name: str) -> str | None:
    m = re.search(r"(\.[A-Za-z0-9]+)$", name or "")
    return KINDS.get(m.group(1).lower()) if m else None


def _check_pdf(data: bytes):
    """Ouverture rapide (index du PDF seulement) : détecte un PDF corrompu ou protégé par mot de passe."""
    try:
        from PyPDF2 import PdfReader
    except Exception:
        return
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise FileReadError("PDF protégé par mot de passe — retirez la protection puis réimportez-le")
        if len(reader.pages) == 0:
            raise FileReadError("PDF sans aucune page")
    except FileReadError:
        raise
    except Exception as exc:
        raise FileReadError(f"PDF corrompu ou illisible ({type(exc).__name__}) — réexportez-le depuis son application d'origine") from exc


def validate_upload(name: str, data: bytes) -> str:
    """Contrôle un fichier téléversé ; renvoie son type ou lève FileReadError."""
    kind = file_kind(name)
    if kind is None and (name or "").lower().endswith((".doc", ".ppt", ".pptx", ".odt")):
        raise FileReadError("ancien format Word/Office non lisible : ouvrez-le et « Enregistrer sous » .docx ou PDF")
    if kind is None:
        raise FileReadError(f"format non pris en charge (formats acceptés : {', '.join(UPLOAD_TYPES)})")
    if not data:
        raise FileReadError("fichier vide (0 octet) — réexportez-le depuis son application d'origine")
    if len(data) > MAX_MB * 1024 * 1024:
        raise FileReadError(f"fichier trop volumineux ({len(data) / 1e6:.0f} Mo, maximum {MAX_MB} Mo) — découpez-le")
    head = data[:8]
    if kind == "pdf" and b"%PDF" not in data[:1024]:
        raise FileReadError("ce fichier n'est pas un vrai PDF (corrompu, ou autre format renommé en .pdf)")
    if kind == "pdf":
        _check_pdf(data)
    ext = name.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm", "docx") and not head.startswith(b"PK"):
        raise FileReadError(f"ce fichier n'est pas un vrai .{ext} (corrompu ou ancien format renommé) — "
                            "ouvrez-le puis « Enregistrer sous » le bon format")
    if ext == "xls" and not (head.startswith(b"\xd0\xcf\x11\xe0") or head.startswith(b"PK")):
        raise FileReadError("ce fichier n'est pas un vrai .xls — ouvrez-le puis enregistrez-le en .xlsx")
    return kind


# =============================================================================
# EXCEL / CSV
# =============================================================================

def _fmt_cell(v) -> tuple[str, str]:
    """Cellule -> (texte, nature) avec nature ∈ {'num', 'txt', ''}. Les vrais nombres sont rendus avec un point."""
    if v is None:
        return "", ""
    if isinstance(v, bool):
        return ("OUI" if v else "NON"), "txt"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:
            return "", ""
        if float(v).is_integer() and abs(v) < 1e15:
            return str(int(v)), "num"
        return ("%.10g" % v), "num"
    if isinstance(v, datetime):
        return (v.strftime("%d/%m/%Y") if v.time() == time(0, 0) else v.strftime("%d/%m/%Y %H:%M")), "txt"
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y"), "txt"
    if isinstance(v, time):
        return v.strftime("%H:%M"), "txt"
    s = str(v).replace(" ", " ").strip()
    return s, "txt"


def _sheet_from_rows(label: str, raw_rows, count_numbers: bool = True) -> SheetPage | None:
    rows, num_typed, str_comma = [], 0, 0
    for raw in raw_rows:
        cells = []
        for v in raw:
            s, nature = _fmt_cell(v)
            if nature == "num" and "." in s:
                num_typed += 1
            elif nature == "txt" and re.fullmatch(r"-?\d[\d ]*,\d+", s):
                str_comma += 1
            cells.append(s)
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            rows.append(cells)
        if len(rows) >= MAX_ROWS:
            break
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    empty_cols = [j for j in range(width) if not any(r[j] for r in rows)]   # colonnes vides -> supprimées
    if empty_cols:
        keep = [j for j in range(width) if j not in empty_cols]
        rows = [[r[j] for j in keep] for r in rows]
    text = "\n".join(" ".join(c for c in r if c) for r in rows)
    dec = "." if count_numbers and num_typed and num_typed >= str_comma else ""
    return SheetPage(label=label, text=text, tables=[rows], decimal=dec, n_rows=len(rows))


def read_excel(name: str, data: bytes) -> list[SheetPage]:
    ext = name.lower().rsplit(".", 1)[-1]
    pages: list[SheetPage] = []
    if ext in ("xlsx", "xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
        except Exception as exc:
            raise FileReadError(f"classeur Excel illisible ou protégé par mot de passe ({type(exc).__name__})") from exc
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            try:
                page = _sheet_from_rows(ws.title, ws.iter_rows(values_only=True))
            except Exception as exc:
                raise FileReadError(f"feuille « {ws.title} » illisible ({type(exc).__name__})") from exc
            if page:
                pages.append(page)
    else:  # .xls ancien format
        try:
            import pandas as pd
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object)
        except ImportError as exc:
            raise FileReadError("ancien format .xls : installez « pip install --user xlrd » ou enregistrez le "
                                "fichier en .xlsx") from exc
        except Exception as exc:
            raise FileReadError(f"classeur .xls illisible ({type(exc).__name__})") from exc
        for title, df in sheets.items():
            page = _sheet_from_rows(str(title), df.where(df.notna(), None).values.tolist())
            if page:
                pages.append(page)
    if not pages:
        raise FileReadError("le classeur ne contient aucune feuille visible avec des données")
    return pages


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_csv(name: str, data: bytes) -> list[SheetPage]:
    text = _decode(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as exc:
        raise FileReadError(f"CSV illisible ({exc})") from exc
    page = _sheet_from_rows(name, rows, count_numbers=False)
    if not page:
        raise FileReadError("le fichier CSV est vide")
    return [page]


def read_text(name: str, data: bytes) -> list[SheetPage]:
    text = _decode(data)
    lines = [ln.rstrip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        raise FileReadError("le fichier texte est vide")
    return [SheetPage(label=name, text="\n".join(lines), n_rows=len(lines))]


# =============================================================================
# WORD (.docx) — lecture directe du XML, sans dépendance
# =============================================================================

def _w_text(node) -> str:
    out = []
    for el in node.iter():
        if el.tag == W + "t":
            out.append(el.text or "")
        elif el.tag == W + "tab":
            out.append("\t")
        elif el.tag == W + "br" and el.get(W + "type") != "page":
            out.append("\n")
        elif el.tag == W + "noBreakHyphen":
            out.append("-")
    return "".join(out)


def _is_page_break(p) -> bool:
    return any(br.get(W + "type") == "page" for br in p.iter(W + "br")) or p.find(f".//{W}pageBreakBefore") is not None


def read_word(name: str, data: bytes) -> list[SheetPage]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        body_xml = zf.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise FileReadError("document Word illisible ou corrompu (pas un .docx valide)") from exc
    try:
        root = ET.fromstring(body_xml)
    except ET.ParseError as exc:
        raise FileReadError("document Word corrompu (XML invalide)") from exc
    body = root.find(W + "body")
    if body is None:
        raise FileReadError("document Word vide")

    head, foot = [], []
    for part in zf.namelist():
        if re.fullmatch(r"word/(header|footer)\d*\.xml", part):
            try:
                txt = " ".join(t for t in (_w_text(p).strip() for p in ET.fromstring(zf.read(part)).iter(W + "p")) if t)
            except ET.ParseError:
                continue
            if txt:
                (head if "header" in part else foot).append(txt)

    pages: list[dict] = [{"lines": [], "tables": []}]
    for child in body:
        if child.tag == W + "p":
            if _is_page_break(child) and pages[-1]["lines"]:
                pages.append({"lines": [], "tables": []})
            txt = _w_text(child).replace("\t", " ").strip()
            for ln in txt.splitlines():
                if ln.strip():
                    pages[-1]["lines"].append(ln.strip())
        elif child.tag == W + "tbl":
            rows = []
            for tr in child.findall(W + "tr"):
                cells = [" ".join(t for t in (_w_text(p).strip() for p in tc.iter(W + "p")) if t)
                         for tc in tr.findall(W + "tc")]
                if any(cells):
                    rows.append(cells)
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                pages[-1]["tables"].append(rows)
                pages[-1]["lines"].extend(" ".join(c for c in r if c) for r in rows)
    if head:
        pages[0]["lines"] = head + pages[0]["lines"]
    if foot:
        pages[-1]["lines"] = pages[-1]["lines"] + foot
    out = [SheetPage(label=f"{name} p.{i}", text="\n".join(p["lines"]), tables=p["tables"], n_rows=len(p["lines"]))
           for i, p in enumerate(pages, 1) if p["lines"]]
    if not out:
        raise FileReadError("le document Word ne contient aucun texte")
    return out


def read_any(name: str, kind: str, data: bytes) -> list[SheetPage]:
    """Point d'entrée : Excel / CSV / Word / texte. (PDF et images sont gérés par le moteur PDF / OCR.)"""
    if kind == "excel":
        return read_excel(name, data)
    if kind == "csv":
        return read_csv(name, data)
    if kind == "word":
        return read_word(name, data)
    if kind == "text":
        return read_text(name, data)
    raise FileReadError(f"format « {kind} » non géré ici")
